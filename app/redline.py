"""Generate evidence-backed visual-redline DOCX drafts (not Word tracked changes)."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor

from app.models import ComplianceFinding, SavedAnalysis


def eligible_findings(analysis: SavedAnalysis, selected_ids: list[str] | None = None) -> list[ComplianceFinding]:
    requested = set(selected_ids) if selected_ids is not None else None
    known = {finding.finding_id for finding in analysis.findings}
    if requested is not None and not requested <= known:
        raise ValueError("One or more finding IDs do not belong to this analysis")
    return [finding for finding in analysis.findings if (requested is None or finding.finding_id in requested) and finding.proposed_revision is not None]


def _add_text(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(text)


def _evidence(document: Document, title: str, evidence) -> None:
    document.add_heading(title, level=3)
    if not evidence:
        document.add_paragraph("No source excerpt is available.")
        return
    for item in evidence:
        paragraph = document.add_paragraph(style="Quote")
        paragraph.add_run(item.text)
        document.add_paragraph(f"Source: {item.document}, page {item.page_number}" + (f", {item.section}" if item.section else ""), style="Caption")


def generate(analysis: SavedAnalysis, findings: list[ComplianceFinding]) -> bytes:
    """Return a valid DOCX package with formatted additions/deletions."""
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Aptos", Pt(10)
    document.add_heading("Contract Compliance Auditor", 0)
    document.add_heading("Visual Redline Draft", 1)
    notice = document.add_paragraph()
    notice_run = notice.add_run("DRAFT FOR LEGAL REVIEW. This document is not an executed amendment or legal opinion.")
    notice_run.bold, notice_run.font.color.rgb = True, RGBColor(0xB9, 0x1C, 0x1C)
    _add_text(document, "Original contract", analysis.contract_document_name)
    _add_text(document, "Policy document", analysis.policy_document_name)
    _add_text(document, "Analysis ID", analysis.analysis_id)
    _add_text(document, "Generated", datetime.now(timezone.utc).isoformat())
    insertions = sum(f.proposed_revision.kind == "insertion" for f in findings if f.proposed_revision)
    replacements = sum(f.proposed_revision.kind == "replacement" for f in findings if f.proposed_revision)
    unresolved = sum(1 for f in analysis.findings if f.status == "ambiguous" or not f.proposed_revision)
    document.add_heading("Executive summary", level=1)
    document.add_paragraph(f"This visual redline draft is based on automated document comparison and requires legal review. Findings considered: {len(findings)}. Proposed insertions: {insertions}. Proposed replacements: {replacements}. Unresolved or ambiguous findings: {unresolved}.")
    document.add_heading("Proposed amendments", level=1)
    reviews = {review.finding_id: review for review in analysis.reviews}
    for finding in findings:
        revision = finding.proposed_revision
        if revision is None:  # Defensive: eligible_findings removes these.
            continue
        document.add_heading(f"{finding.finding_id} — {revision.kind.title()}", level=2)
        _add_text(document, "Compliance status", finding.status)
        _add_text(document, "Risk level", finding.risk_level)
        _add_text(document, "Requirement", finding.policy_evidence[0].text if finding.policy_evidence else finding.requirement_id)
        review = reviews.get(finding.finding_id)
        _add_text(document, "Review status", f"{review.decision} at {review.reviewed_at}" if review else "Unreviewed AI-generated proposal")
        _add_text(document, "Reason", revision.rationale)
        if revision.kind == "replacement" and finding.contract_evidence:
            document.add_paragraph("Existing contract wording (proposed deletion):")
            old = document.add_paragraph().add_run(finding.contract_evidence[0].text)
            old.font.strike = True
            old.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            document.add_paragraph("Proposed replacement wording (proposed addition):")
        elif revision.kind == "insertion":
            document.add_paragraph("Proposed new clause (insertion location requires legal review):")
        else:
            document.add_paragraph("No reliable existing clause text is available; this is a recommendation requiring legal review.")
        addition = document.add_paragraph().add_run(revision.text)
        addition.font.color.rgb = RGBColor(0x16, 0x7A, 0x3A)
        addition.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
        _evidence(document, "Policy evidence", finding.policy_evidence)
        _evidence(document, "Contract evidence", finding.contract_evidence)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
