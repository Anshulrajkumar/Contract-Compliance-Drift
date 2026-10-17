"""DOCX visual-redline generation and artifact API tests."""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.shared import RGBColor
from fastapi.testclient import TestClient

from app import database, registry, store
from app.models import (AnalysisSummary, ComparisonAnalysis, ComplianceFinding,
    DocumentInfo, EvidenceExcerpt, ProposedRevision)


@pytest.fixture(autouse=True)
def isolated_state(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "analyses.db"))
    monkeypatch.setenv("ARTIFACT_DIR", str(tmp_path / "redlines"))
    registry.reset_registry(); store.clear(); database.initialize()
    yield
    registry.reset_registry(); store.clear()


@pytest.fixture()
def client():
    from app.main import app
    with TestClient(app) as instance:
        yield instance


def seeded_analysis(kind="replacement", with_revision=True) -> str:
    analysis_id = "redline-analysis"
    database.create_pending(analysis_id, "policy", "contract", "Policy.pdf", "Vendor MSA.pdf")
    database.mark_running(analysis_id)
    policy = EvidenceExcerpt(text="Supplier shall protect confidential data.", document="Policy.pdf", page_number=3, section="Confidentiality")
    contract = EvidenceExcerpt(text="Supplier may disclose Customer information.", document="Vendor MSA.pdf", page_number=2, section="8.1")
    revision = ProposedRevision(kind=kind, text="Supplier shall protect Customer Confidential Information.", rationale="Aligns the contract with the cited policy.") if with_revision else None
    finding = ComplianceFinding(finding_id="finding-001", requirement_id="req-001", status="potentially_non_compliant", risk_level="high", summary="Confidentiality protection needs revision.", reasoning_summary="The contract clause is less protective.", policy_evidence=[policy], contract_evidence=[contract], search_scope="All contract pages searched.", proposed_revision=revision, confidence=.9, review_required=True)
    analysis = ComparisonAnalysis(analysis_id=analysis_id, policy_document=DocumentInfo(document_id="policy", filename="Policy.pdf", size_bytes=1), contract_document=DocumentInfo(document_id="contract", filename="Vendor MSA.pdf", size_bytes=1), requirements_analyzed=1, findings=[finding], summary=AnalysisSummary(potentially_non_compliant=1), warnings=[])
    database.save_completed(analysis)
    return analysis_id


def test_replacement_docx_has_visual_addition_deletion_and_evidence(client):
    analysis_id = seeded_analysis()
    response = client.post(f"/api/analyses/{analysis_id}/redlines", json={})
    assert response.status_code == 201
    artifact = response.json()
    download = client.get(artifact["download_url"])
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert download.content[:2] == b"PK"  # DOCX is a zip package.
    document = Document(BytesIO(download.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "draft for legal review" in text.lower()
    assert "Policy.pdf, page 3" in text
    assert any(run.font.strike for p in document.paragraphs for run in p.runs if "may disclose" in run.text)
    assert any(run.font.color.rgb == RGBColor(0x16, 0x7A, 0x3A) for p in document.paragraphs for run in p.runs if "shall protect" in run.text)


def test_insertion_is_labelled_without_invented_clause_location(client):
    analysis_id = seeded_analysis(kind="insertion")
    artifact = client.post(f"/api/analyses/{analysis_id}/redlines", json={}).json()
    document = Document(BytesIO(client.get(artifact["download_url"]).content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Proposed new clause" in text
    assert "insertion location requires legal review" in text


def test_invalid_ids_and_missing_revisions_are_safe(client):
    assert client.post("/api/analyses/nope/redlines", json={}).status_code == 404
    analysis_id = seeded_analysis(with_revision=False)
    assert client.post(f"/api/analyses/{analysis_id}/redlines", json={}).status_code == 422
    assert client.get("/api/redlines/nope/download").status_code == 404


def test_selected_finding_validation_and_artifact_listing(client):
    analysis_id = seeded_analysis()
    assert client.post(f"/api/analyses/{analysis_id}/redlines", json={"finding_ids": ["other"]}).status_code == 422
    artifact = client.post(f"/api/analyses/{analysis_id}/redlines", json={"finding_ids": ["finding-001"]}).json()
    listed = client.get(f"/api/analyses/{analysis_id}/redlines").json()
    assert listed[0]["artifact_id"] == artifact["artifact_id"]
